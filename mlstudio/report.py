"""Word (.docx) report of the whole analysis, with charts and a ready-to-use methods paragraph."""
from __future__ import annotations

import datetime as _dt
import io

import numpy as np
import pandas as pd

from . import APP_NAME, APP_VERSION, INSTITUTION
from .models import MODELS, describe_params, model_name
from .trainer import METRIC_HELP, model_weights, permutation_importance_table, quality_word

BRAND = (0x7B, 0x1C, 0x2E)
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]


def _sk_version():
    import sklearn
    return sklearn.__version__


def methods_text(ss) -> str:
    task, target, feats = ss["task"], ss["target"], ss["features"]
    sp = ss["split"]
    res = ss["results"]
    labels = sp["labels"]
    n_tr, n_te = len(sp["y_train"]), len(sp["y_test"])
    n = n_tr + n_te
    binary = task == "classification" and labels and len(labels) == 2
    parts = []
    parts.append(f"Data from {n} patients with a recorded outcome were analysed.")
    nlog = len([e for e in ss.get("clean_log", []) if e.get("code")])
    if nlog:
        parts.append(f"Before modelling, the dataset was checked for data-entry problems and {nlog} cleaning step(s) were "
                     "applied (listed in the supplementary cleaning log).")
    if binary:
        parts.append(f"The outcome was {target} (positive class: '{labels[1]}'), modelled as a binary classification problem.")
    elif task == "classification":
        parts.append(f"The outcome was {target} with {len(labels)} categories ({', '.join(map(str, labels))}), modelled as a "
                     "multi-class classification problem.")
    else:
        parts.append(f"The outcome was {target}, a continuous variable, modelled as a regression problem.")
    parts.append(f"{len(feats)} candidate predictors were used: {', '.join(feats)}.")
    strat = " stratified by outcome" if (task == "classification" and sp.get("stratify", True)) else ""
    parts.append(f"Patients were randomly split{strat} into a training set (n = {n_tr}, {100 - sp['test_pct']}%) and a held-out "
                 f"test set (n = {n_te}, {sp['test_pct']}%; random seed {sp['seed']}).")
    parts.append("Within a single modelling pipeline fitted on the training data only, missing values were imputed with the "
                 "median (continuous predictors) or the most frequent category (categorical predictors), categorical "
                 "predictors were one-hot encoded, and continuous predictors were standardised for algorithms that are "
                 "sensitive to scale.")
    names = [r["name"] for r in res.values()]
    parts.append(f"The following algorithms were compared: {', '.join(names)}.")
    tuned = [r["name"] for r in res.values() if r["mode"] in ("grid", "random")]
    cvk = f"{sp['folds']}-fold {'stratified ' if task == 'classification' else ''}cross-validation"
    if tuned:
        parts.append(f"Hyperparameters of {', '.join(tuned)} were tuned by "
                     f"{'grid' if any(r['mode'] == 'grid' for r in res.values()) else 'random'} search using {cvk} on the "
                     f"training set, with {sp['metric']} as the selection criterion; the other models used pre-specified settings.")
    else:
        parts.append(f"Model performance in the training set was estimated with {cvk} ({sp['metric']}).")
    if sp.get("balance") and task == "classification":
        parts.append("Class imbalance was addressed with balanced class weights.")
    parts.append("Each final model was refitted on the complete training set and evaluated once on the test set.")
    if binary:
        parts.append("Discrimination was assessed with the area under the receiver operating characteristic curve (AUC); "
                     "sensitivity, specificity, positive and negative predictive values were calculated at a probability "
                     "threshold of 0.5. Calibration was assessed with a calibration plot and the Brier score.")
    elif task == "classification":
        parts.append("Performance was summarised by accuracy, balanced accuracy, macro-averaged F1 and the one-vs-rest AUC, "
                     "with a confusion matrix.")
    else:
        parts.append("Performance was summarised by R², mean absolute error (MAE) and root mean squared error (RMSE).")
    parts.append("95% confidence intervals for test-set metrics were obtained by bootstrap resampling of the test set "
                 "(400 resamples, percentile method). Predictor importance was estimated by permutation importance on the "
                 "test set.")
    parts.append(f"Analyses were performed with {APP_NAME} version {APP_VERSION} (Python, scikit-learn {_sk_version()}).")
    return " ".join(parts)


def results_text(ss) -> str:
    task = ss["task"]
    res = ss["results"]
    labels = ss["split"]["labels"]
    binary = task == "classification" and labels and len(labels) == 2
    main = "AUC" if task == "classification" else "R2"
    best_k = max(res, key=lambda k: np.nan_to_num(res[k]["test"].get(main, np.nan), nan=-9))
    r = res[best_k]
    t = r["test"]
    lo, hi = r["ci"].get(main, (np.nan, np.nan))
    if binary:
        return (f"On the held-out test set, {r['name']} showed the highest discrimination (AUC {t['AUC']:.3f}, 95% CI "
                f"{lo:.3f}-{hi:.3f}), with a sensitivity of {t['Sensitivity'] * 100:.1f}% and a specificity of "
                f"{t['Specificity'] * 100:.1f}% at a threshold of 0.5 (PPV {t['PPV'] * 100:.1f}%, NPV {t['NPV'] * 100:.1f}%).")
    if task == "classification":
        return (f"On the held-out test set, {r['name']} performed best (one-vs-rest AUC {t['AUC']:.3f}, 95% CI {lo:.3f}-{hi:.3f}; "
                f"accuracy {t['Accuracy'] * 100:.1f}%, balanced accuracy {t['Balanced accuracy'] * 100:.1f}%).")
    mlo, mhi = r["ci"].get("MAE", (np.nan, np.nan))
    return (f"On the held-out test set, {r['name']} explained the most variation in {ss['target']} (R² {t['R2']:.3f}, 95% CI "
            f"{lo:.3f}-{hi:.3f}); the mean absolute error was {t['MAE']:.2f} (95% CI {mlo:.2f}-{mhi:.2f}).")


# ---------------------------------------------------------------------------
# matplotlib charts for the report
# ---------------------------------------------------------------------------
def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False, "axes.edgecolor": "#898781",
                         "axes.titleweight": "bold", "axes.titlelocation": "left", "axes.titlesize": 10})
    return plt


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    buf.seek(0)
    import matplotlib.pyplot as plt
    plt.close(fig)
    return buf


def fig_roc(ss):
    from sklearn.metrics import roc_curve
    plt = _mpl()
    yte = ss["split"]["y_test"]
    fig, ax = plt.subplots(figsize=(4.6, 4.4))
    ax.plot([0, 1], [0, 1], ls=":", color="#c3c2b7", lw=1.2)
    for i, r in enumerate(ss["results"].values()):
        fpr, tpr, _ = roc_curve(yte, r["proba_test"])
        ax.plot(fpr, tpr, lw=2, color=SERIES[i % 8], label=f"{r['name']} (AUC {r['test']['AUC']:.3f})")
    ax.set_xlabel("1 - Specificity")
    ax.set_ylabel("Sensitivity")
    ax.set_title("ROC curves on the test set")
    ax.legend(loc="lower right", fontsize=7.5, frameon=False)
    return _png(fig)


def fig_calibration(ss):
    from sklearn.calibration import calibration_curve
    plt = _mpl()
    yte = ss["split"]["y_test"]
    fig, ax = plt.subplots(figsize=(4.6, 4.4))
    ax.plot([0, 1], [0, 1], ls=":", color="#c3c2b7", lw=1.2)
    nb = int(min(10, max(3, len(yte) // 15)))
    for i, r in enumerate(ss["results"].values()):
        fp, mp = calibration_curve(yte, r["proba_test"], n_bins=nb, strategy="quantile")
        ax.plot(mp, fp, marker="o", lw=1.6, ms=4, color=SERIES[i % 8], label=r["name"])
    ax.set_xlabel("Predicted probability")
    ax.set_ylabel("Observed proportion")
    ax.set_title("Calibration on the test set")
    ax.legend(fontsize=7.5, frameon=False)
    return _png(fig)


def fig_confusion(cm, labels, title):
    plt = _mpl()
    cm = np.asarray(cm)
    fig, ax = plt.subplots(figsize=(1.6 + 1.0 * len(labels), 1.3 + 0.9 * len(labels)))
    rowpct = cm / np.maximum(cm.sum(axis=1, keepdims=True), 1)
    ax.imshow(rowpct, cmap="Blues", vmin=0, vmax=1)
    for i in range(cm.shape[0]):
        for j in range(cm.shape[1]):
            ax.text(j, i, f"{cm[i, j]}\n({rowpct[i, j] * 100:.0f}%)", ha="center", va="center", fontsize=8,
                    color="white" if rowpct[i, j] > 0.6 else "#0b0b0b")
    ax.set_xticks(range(len(labels)), [f"Pred. {l}" for l in labels], rotation=20, ha="right", fontsize=8)
    ax.set_yticks(range(len(labels)), [f"True {l}" for l in labels], fontsize=8)
    ax.set_title(title)
    for s in ax.spines.values():
        s.set_visible(False)
    return _png(fig)


def fig_importance(perm: pd.DataFrame, title):
    plt = _mpl()
    d = perm.head(15).iloc[::-1]
    fig, ax = plt.subplots(figsize=(5.6, 0.32 * len(d) + 1.1))
    ax.barh(d["Factor"].astype(str), d["Importance"], xerr=d["SD"], color=SERIES[0], ecolor="#898781", height=0.6)
    ax.set_xlabel(f"Drop in {perm['metric'].iloc[0]} when scrambled")
    ax.set_title(title)
    return _png(fig)


def fig_pred_actual(y, pred, title):
    plt = _mpl()
    fig, ax = plt.subplots(figsize=(4.4, 4.2))
    lo, hi = min(np.min(y), np.min(pred)), max(np.max(y), np.max(pred))
    ax.plot([lo, hi], [lo, hi], ls=":", color="#c3c2b7")
    ax.scatter(y, pred, s=14, alpha=0.7, color=SERIES[0], edgecolors="white", linewidths=0.5)
    ax.set_xlabel("Actual")
    ax.set_ylabel("Predicted")
    ax.set_title(title)
    return _png(fig)


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
def _table(doc, df: pd.DataFrame, digits=3, widths=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(c))
        run.bold = True
        run.font.size = Pt(8.5)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), "F3EDEE")
        cell._tc.get_or_add_tcPr().append(sh)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            if isinstance(v, (float, np.floating)):
                txt = "-" if not np.isfinite(v) else f"{v:.{digits}f}"
            else:
                txt = str(v)
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(txt)
            r.font.size = Pt(8.5)
    doc.add_paragraph()
    return t


def build_docx(ss) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    task, target = ss["task"], ss["target"]
    sp = ss["split"]
    res = ss["results"]
    labels = sp["labels"]
    binary = task == "classification" and labels and len(labels) == 2

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for lvl in (1, 2):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.color.rgb = RGBColor(*BRAND)
        hs.font.name = "Calibri"

    try:
        from .ui import ASSETS
        logo = ASSETS / "mamc_logo.png"
        if logo.exists():
            doc.add_picture(str(logo), width=Cm(2.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass
    h = doc.add_heading("Machine-learning analysis report", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"{APP_NAME} v{APP_VERSION} · {INSTITUTION}\n"
                          f"Dataset: {ss.get('data_name')} · Date: {_dt.date.today().strftime('%d %B %Y')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. data
    doc.add_heading("1. Data", 1)
    df = ss["df"]
    raw = ss["df_raw"]
    doc.add_paragraph(f"Original file: {raw.shape[0]} rows x {raw.shape[1]} columns. After cleaning: {df.shape[0]} rows x "
                      f"{df.shape[1]} columns. Patients with an outcome used for modelling: {len(sp['y_train']) + len(sp['y_test'])}.")
    log = ss.get("clean_log", [])
    if log:
        doc.add_heading("Cleaning log", 2)
        _table(doc, pd.DataFrame({"Step": range(1, len(log) + 1), "What was done": [e["text"] for e in log]}))
    doc.add_heading("Outcome and predictors", 2)
    if task == "classification":
        y_all = np.r_[sp["y_train"], sp["y_test"]]
        if binary:
            y_all = np.where(y_all == 1, labels[1], labels[0])
        vc = pd.Series(y_all).value_counts()
        doc.add_paragraph(f"Outcome: {target} ({'binary' if binary else 'multi-class'} classification" +
                          (f"; positive class '{labels[1]}'" if binary else "") + "). Distribution: " +
                          "; ".join(f"{k}: {v} ({v / vc.sum() * 100:.1f}%)" for k, v in vc.items()) + ".")
    else:
        y_all = np.r_[sp["y_train"], sp["y_test"]]
        doc.add_paragraph(f"Outcome: {target} (regression). Mean {np.mean(y_all):.2f}, SD {np.std(y_all, ddof=1):.2f}, "
                          f"median {np.median(y_all):.2f}.")
    doc.add_paragraph("Predictors (" + str(len(ss["features"])) + "): " + ", ".join(ss["features"]) + ".")

    # 2. methods
    doc.add_heading("2. Methods (ready to adapt for your thesis)", 1)
    doc.add_paragraph(methods_text(ss))
    rows = []
    for k, r in res.items():
        rows.append({"Model": r["name"], "How settings were chosen": {"default": "Pre-specified starting values",
                     "manual": "Chosen by the user", "grid": "Grid search", "random": "Random search"}[r["mode"]],
                     "Final settings": "; ".join(describe_params(k, r["params"])) or "none"})
    _table(doc, pd.DataFrame(rows))

    # 3. results
    doc.add_heading("3. Results", 1)
    doc.add_paragraph(results_text(ss))
    if binary:
        keys = ["AUC", "Sensitivity", "Specificity", "PPV", "NPV", "Balanced accuracy", "Brier score"]
    elif task == "classification":
        keys = ["AUC", "Accuracy", "Balanced accuracy", "F1 (macro)"]
    else:
        keys = ["R2", "MAE", "RMSE"]
    m = sp["metric"]
    rows = []
    for k, r in res.items():
        row = {"Model": r["name"], f"CV {m} (mean ± SD)": f"{r['cv'][m]['mean']:.3f} ± {r['cv'][m]['sd']:.3f}"}
        for key in keys:
            v = r["test"].get(key, np.nan)
            lo, hi = r["ci"].get(key, (np.nan, np.nan))
            row[f"Test {key}"] = f"{v:.3f}" + (f" ({lo:.3f}-{hi:.3f})" if np.isfinite(lo) else "") if np.isfinite(v) else "-"
        rows.append(row)
    doc.add_paragraph("Table: performance in cross-validation (training set) and on the held-out test set "
                      "(95% bootstrap confidence intervals in brackets).").runs[0].italic = True
    _table(doc, pd.DataFrame(rows))

    main = "AUC" if task == "classification" else "R2"
    best_k = max(res, key=lambda k: np.nan_to_num(res[k]["test"].get(main, np.nan), nan=-9))
    best = res[best_k]
    if binary:
        doc.add_picture(fig_roc(ss), width=Cm(8.0))
        doc.add_picture(fig_calibration(ss), width=Cm(8.0))
        t = best["test"]
        doc.add_picture(fig_confusion([[t["TN"], t["FP"]], [t["FN"], t["TP"]]], labels, f"Confusion matrix - {best['name']}"),
                        width=Cm(7.0))
    elif task == "classification":
        doc.add_picture(fig_confusion(best["test"]["confusion"], best["classes"], f"Confusion matrix - {best['name']}"),
                        width=Cm(10.0))
        _table(doc, best["test"]["per_class"])
    else:
        doc.add_picture(fig_pred_actual(sp["y_test"], best["pred_test"], f"Predicted vs actual - {best['name']}"), width=Cm(8.0))

    doc.add_heading("Model fit (over- / under-fitting)", 2)
    for k, r in res.items():
        doc.add_paragraph(f"{r['name']}: {r['verdict']['label']}. {r['verdict']['text']}", style="List Bullet")

    doc.add_heading("Important predictors", 2)
    perm = best.get("perm")
    if perm is None:
        perm = permutation_importance_table(best, sp["X_test"], sp["y_test"], task, labels, seed=sp["seed"])
    doc.add_picture(fig_importance(perm, f"Permutation importance - {best['name']}"), width=Cm(12))
    w = model_weights(best, task, labels)
    if w is not None and "Odds ratio" in w:
        doc.add_paragraph("Odds ratios (per 1 SD for continuous predictors) from the logistic regression:")
        _table(doc, w[["Term", "Odds ratio", "Direction"]].head(20), digits=2)

    # 4. interpretation & limitations
    doc.add_heading("4. Interpretation and limitations", 1)
    for txt in [
        "The test set comes from the same data source as the training set (internal validation). Performance in other "
        "hospitals, populations or time periods may be lower; external validation is needed before any clinical use.",
        f"Sample size: {len(sp['y_train']) + len(sp['y_test'])} patients. With small datasets, results vary with the random "
        "split; confidence intervals show this uncertainty.",
        "Missing values were imputed with simple methods (median / most frequent). If data are not missing at random, "
        "this can bias the model; multiple imputation may be preferable for inference.",
        "Permutation importance shows what the model relies on, not causal effects.",
        "Report the model following the TRIPOD+AI statement (Collins et al., BMJ 2024;385:e078378).",
    ]:
        doc.add_paragraph(txt, style="List Bullet")

    doc.add_heading("Metric glossary", 2)
    for k in keys:
        if k in METRIC_HELP:
            doc.add_paragraph(f"{k}: {METRIC_HELP[k]}", style="List Bullet")

    doc.add_heading("References", 2)
    for ref in [
        "Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction "
        "models that use regression or machine learning methods. BMJ. 2024;385:e078378.",
        "Riley RD, Ensor J, Snell KIE, et al. Calculating the sample size required for developing a clinical prediction "
        "model. BMJ. 2020;368:m441.",
        "Pedregosa F, Varoquaux G, Gramfort A, et al. Scikit-learn: machine learning in Python. J Mach Learn Res. "
        "2011;12:2825-30.",
    ]:
        doc.add_paragraph(ref, style="List Number")
    doc.add_paragraph()
    f = doc.add_paragraph(f"Generated by {APP_NAME} v{APP_VERSION}, Maulana Azad Medical College, New Delhi. "
                          "For learning and research; not for clinical decisions on individual patients.")
    f.runs[0].italic = True
    f.runs[0].font.size = Pt(8.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
